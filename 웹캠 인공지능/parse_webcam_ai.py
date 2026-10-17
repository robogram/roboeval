# -*- coding: utf-8 -*-
import openpyxl
import json
import os
import sys
import io

# UTF-8 출력 설정
sys.stdout = io.TextIOWrapper(sys.stdout.detach(), encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.detach(), encoding='utf-8')

def parse_webcam_ai_questions():
    # 파일 경로
    current_dir = os.path.dirname(os.path.abspath(__file__))
    # 루트 디렉토리 기준 또는 현재 폴더 기준 엑셀 파일 탐색
    possible_paths = [
        os.path.join(current_dir, "웹캠_AI_단원별_4지선다_100문제.xlsx"),
        os.path.join(current_dir, "..", "웹캠_AI_단원별_4지선다_100문제.xlsx"),
        r"E:\AI 개발 협업\robogram 성취도평가\roboeval\웹캠_AI_단원별_4지선다_100문제.xlsx"
    ]
    
    excel_path = None
    for p in possible_paths:
        if os.path.exists(p):
            excel_path = p
            break
            
    if not excel_path:
        print("오류: '웹캠_AI_단원별_4지선다_100문제.xlsx' 파일을 찾을 수 없습니다.")
        return None

    print(f"엑셀 파일 로딩 중: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 1. '정답·해설' 시트 파싱
    if '정답·해설' not in wb.sheetnames:
        print("오류: '정답·해설' 시트가 존재하지 않습니다.")
        return None

    ans_sheet = wb['정답·해설']
    ans_map = {}
    for r in range(7, ans_sheet.max_row + 1):
        q_id = ans_sheet.cell(r, 1).value
        if not q_id:
            continue
        q_id_str = str(q_id).strip()
        unit = ans_sheet.cell(r, 2).value
        q_num = ans_sheet.cell(r, 3).value
        diff = ans_sheet.cell(r, 4).value
        ans_num = ans_sheet.cell(r, 5).value
        ans_text = ans_sheet.cell(r, 6).value
        explanation = ans_sheet.cell(r, 7).value
        reference = ans_sheet.cell(r, 8).value
        
        ans_map[q_id_str] = {
            "unit": str(unit).strip() if unit else "",
            "order": int(q_num) if q_num is not None else 0,
            "diff": str(diff).strip() if diff else "기초",
            "ans_num": int(ans_num) if ans_num is not None else 1,
            "ans_text": str(ans_text).strip() if ans_text else "",
            "explanation": str(explanation).strip() if explanation else "",
            "reference": str(reference).strip() if reference else ""
        }

    print(f"정답·해설 데이터 파싱 완료: {len(ans_map)}개 문항")

    # 2. 단원별 문제 시트 파싱
    unit_sheets = [
        (1, '1단원 얼굴인식', '얼굴인식'),
        (2, '2단원 사물인식', '사물인식'),
        (3, '3단원 AI 페르소나', 'AI 페르소나'),
        (4, '4단원 캡스톤 통합', '캡스톤 통합'),
        (5, '5단원 생성형 AI', '생성형 AI')
    ]

    all_questions = []
    global_q_idx = 1

    for day_val, sname, domain_name in unit_sheets:
        if sname not in wb.sheetnames:
            print(f"경고: 시트 '{sname}'을 찾을 수 없습니다.")
            continue
            
        sheet = wb[sname]
        unit_count = 0
        
        for r in range(7, sheet.max_row + 1):
            q_id = sheet.cell(r, 1).value
            if not q_id:
                continue
            q_id_str = str(q_id).strip()
            diff_val = sheet.cell(r, 2).value
            question_text = sheet.cell(r, 3).value
            c1 = sheet.cell(r, 4).value
            c2 = sheet.cell(r, 5).value
            c3 = sheet.cell(r, 6).value
            c4 = sheet.cell(r, 7).value

            choices = [
                str(c1).strip() if c1 is not None else "",
                str(c2).strip() if c2 is not None else "",
                str(c3).strip() if c3 is not None else "",
                str(c4).strip() if c4 is not None else ""
            ]

            ans_info = ans_map.get(q_id_str, {})
            # 정답 번호 1~4를 0-based 인덱스(0~3)로 변환 (AGENTS.md 규칙)
            ans_1based = ans_info.get("ans_num", 1)
            ans_0based = int(ans_1based) - 1

            # 검증: 선택지와 정답 텍스트 비교
            if 0 <= ans_0based < len(choices):
                actual_choice_text = choices[ans_0based]
                expected_ans_text = ans_info.get("ans_text", "")
                if actual_choice_text != expected_ans_text:
                    print(f"주의: 문항 {q_id_str} 보기[{ans_0based+1}]과 정답해설 텍스트 불일치 가능성")

            order_val = ans_info.get("order", unit_count + 1)
            item_id = f"WEBCAM-AI-Q{global_q_idx:03d}"

            q_obj = {
                "item_id": item_id,
                "raw_id": q_id_str,
                "day": day_val,
                "order": order_val,
                "type": "객관식",
                "content": str(question_text).strip(),
                "choices": choices,
                "answer": ans_0based, # 0-based
                "domain": domain_name,
                "diff": str(diff_val).strip() if diff_val else ans_info.get("diff", "기초"),
                "explanation": ans_info.get("explanation", ""),
                "reference": ans_info.get("reference", "")
            }
            all_questions.append(q_obj)
            global_q_idx += 1
            unit_count += 1

        print(f"[{sname}] 파싱 완료: {unit_count}문항 (Day {day_val})")

    print(f"\n최종 파싱 완료 문항 수: {len(all_questions)}개")

    # JSON 저장
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_json_path = os.path.join(output_dir, "questions_parsed.json")
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(all_questions, f, ensure_ascii=False, indent=2)
    print(f"문항 JSON 저장 성공: {output_json_path}")

    return all_questions

if __name__ == "__main__":
    parse_webcam_ai_questions()
