# -*- coding: utf-8 -*-
import openpyxl
import json
import os
import re

def parse_and_inject_pattern():
    excel_path = r"E:\파이썬 문제만들기\고등학생용 문제만들기\고등학생용_파이썬_단원별_4지선다_문제은행.xlsx"
    
    if not os.path.exists(excel_path):
        print(f"오류: 고등 패턴 파이썬 엑셀 문제집 파일을 찾을 수 없습니다. 경로: {excel_path}")
        return

    print(f"고등 엑셀 파일 로딩 중: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # '프로그래밍실습 문제은행' 시트 선택
    sheet_name = '프로그래밍실습 문제은행'
    if sheet_name not in wb.sheetnames:
        print(f"오류: '{sheet_name}' 시트가 존재하지 않습니다.")
        return
        
    sheet = wb[sheet_name]
    print(f"활성화된 시트명: {sheet.title}")

    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]

    # 컬럼 인덱스 매핑
    try:
        idx_day = header.index('단원번호')
        idx_order = header.index('문항번호')
        idx_content = header.index('문제')
        idx_opt1 = header.index('보기①')
        idx_opt2 = header.index('보기②')
        idx_opt3 = header.index('보기③')
        idx_opt4 = header.index('보기④')
        idx_ans_num = header.index('정답')
    except ValueError as e:
        print(f"오류: 엑셀 파일의 헤더 이름이 올바르지 않습니다. 필수 컬럼들을 확인하세요: {e}")
        return

    questions = []
    
    for idx, r in enumerate(rows[1:], start=1):
        if r[idx_day] is None or r[idx_content] is None:
            continue
        
        try:
            day_val = int(r[idx_day])
            order_val = int(r[idx_order]) if r[idx_order] is not None else idx
        except (ValueError, TypeError):
            continue

        content = str(r[idx_content]).strip()
        choices = [
            str(r[idx_opt1]).strip() if r[idx_opt1] is not None else "",
            str(r[idx_opt2]).strip() if r[idx_opt2] is not None else "",
            str(r[idx_opt3]).strip() if r[idx_opt3] is not None else "",
            str(r[idx_opt4]).strip() if r[idx_opt4] is not None else ""
        ]
        
        # 빈 선택지 제거
        choices = [c for c in choices if c != ""]
        
        # 정답 번호 0-based index 변환
        ans_num_str = str(r[idx_ans_num]).strip()
        try:
            ans_val = int(float(ans_num_str)) - 1
        except (ValueError, TypeError):
            ans_val = 0
            print(f"경고: 문항번호 {order_val} (Day {day_val})의 정답 번호 '{ans_num_str}' 파싱 실패. 기본값 0으로 설정합니다.")

        questions.append({
            "item_id": f"ROBOME-AVATAR-PATTERN-Q{idx:03d}",
            "day": day_val,
            "order": order_val,
            "type": "객관식",
            "content": content,
            "choices": choices,
            "answer": ans_val
        })

    print(f"성공적으로 {len(questions)}개의 문항을 파싱했습니다.")

    # index.html 파일 로드 및 주입
    html_path = r"e:\AI 개발 협업\robogram 성취도평가\roboeval\index.html"
    if not os.path.exists(html_path):
        print("오류: index.html 파일을 찾을 수 없습니다.")
        return

    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    pattern_js_str = "const MASTER_AVATAR_PATTERN_QUESTIONS = " + json.dumps(questions, ensure_ascii=False, indent=2) + ";"

    start_marker = "const MASTER_AVATAR_PATTERN_QUESTIONS = ["
    end_marker = "const MASTER_FUSION_QUESTIONS = ["

    start_idx = html_content.find(start_marker)
    end_idx = html_content.find(end_marker)

    if start_idx != -1 and end_idx != -1:
        html_content = html_content[:start_idx] + pattern_js_str + "\n\n" + html_content[end_idx:]
        print("index.html 내의 MASTER_AVATAR_PATTERN_QUESTIONS 상수를 성공적으로 교체했습니다!")
    else:
        print("오류: index.html 내에서 MASTER_AVATAR_PATTERN_QUESTIONS 선언 영역을 찾을 수 없습니다.")
        return

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print("index.html 파일 저장 완료!")

if __name__ == "__main__":
    parse_and_inject_pattern()
